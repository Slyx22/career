"""
Build app/data/sa_occupations.json: South African job titles (OFO 2021)
linked to the O*NET-based careers the engine scores against.

Chain:  OFO occupation (6-digit)  ->  OFO unit group (4-digit)
        ->  ISCO-08 unit group     (matched by title; OFO is built on ISCO-08
                                     but renumbers some groups, e.g. trades)
        ->  2010 SOC codes         (BLS ISCO-08 <-> 2010 SOC crosswalk)
        ->  2018 SOC codes         (BLS 2010 -> 2018 SOC crosswalk)
        ->  O*NET careers          (career files' onet.soc_code)
Within those candidate careers, the one whose name best matches the SA
title (and its OFO specialisations) is chosen. If the chain gives no
candidates, a strong direct title match is used instead, otherwise the SA
title is left out rather than guessed.

Source files (all free) live in app/data/sources/ and are fetched by the
"Fetch South African occupation sources" GitHub workflow:
  ofo_2021.xlsx                     - OFO 2021 (via HWSETA)
  isco08_soc2010_crosswalk.xls      - U.S. BLS
  soc2010_to_soc2018_crosswalk.xlsx - U.S. BLS

Run:  python scripts/build_sa_occupations.py   (needs openpyxl and xlrd)
"""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Set, Tuple

ROOT = Path(__file__).resolve().parent.parent
SOURCES = ROOT / "app" / "data" / "sources"
CAREERS_DIR = ROOT / "app" / "data" / "career_models"
OUT = ROOT / "app" / "data" / "sa_occupations.json"

STOP = {
    "and", "or", "the", "of", "a", "an", "n", "e", "c", "not", "elsewhere", "classified",
    "related", "other", "in", "nec", "all", "for", "to", "general", "except", "including",
}
# Words that differ between SA/UK and US spelling or naming.
SYNONYMS = {
    "paediatric": "pediatric", "paediatrics": "pediatric", "anaesthetist": "anesthesiologist",
    "anaesthesiologist": "anesthesiologist", "orthopaedic": "orthopedic", "gynaecologist": "gynecologist",
    "labour": "labor", "organisation": "organization", "programme": "program", "centre": "center",
    "draughtsperson": "drafter", "draughtsman": "drafter", "lorry": "truck", "chemist": "chemist",
    "attorney": "lawyer", "advocate": "lawyer", "educator": "teacher", "practitioner": "practitioner",
    "salesperson": "sales", "shopkeeper": "retail",
}


def toks(text: str) -> Set[str]:
    words = re.findall(r"[a-z]+", (text or "").lower())
    out = set()
    for w in words:
        if w in STOP:
            continue
        w = SYNONYMS.get(w, w)
        if len(w) > 4 and w.endswith("ies"):
            w = w[:-3] + "y"
        elif len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
            w = w[:-1]
        out.add(w)
    return out


def sim(a: str, b: str) -> float:
    A, B = toks(a), toks(b)
    if not A or not B:
        return 0.0
    return len(A & B) / len(A | B)


def overlap(a: str, b: str) -> float:
    """Share of the SA title's words found in the career name."""
    A, B = toks(a), toks(b)
    return len(A & B) / len(A) if A else 0.0


def _ofo_rows(sheet) -> List[Tuple[str, str]]:
    rows = []
    for r in list(sheet.iter_rows(values_only=True))[3:]:
        if r and r[0] and r[1] and str(r[0]).startswith("2021-"):
            rows.append((str(r[0])[5:].strip(), str(r[1]).strip()))
    return rows


def load_ofo():
    import openpyxl

    wb = openpyxl.load_workbook(SOURCES / "ofo_2021.xlsx", read_only=True)
    units = dict(_ofo_rows(wb["Unit Groups"]))
    occupations = _ofo_rows(wb["Occupations"])
    alt: Dict[str, List[str]] = defaultdict(list)
    for sheet in ("Specialisations", "Alternate Title"):
        for code, title in _ofo_rows(wb[sheet]):
            for part in re.split(r"\s*/\s*", title):
                part = part.strip()
                if part and part.lower() != "test":
                    alt[code].append(part)
    return units, occupations, alt


def load_isco_to_soc2010():
    import xlrd

    sheet = xlrd.open_workbook(SOURCES / "isco08_soc2010_crosswalk.xls").sheet_by_name("ISCO-08 to 2010 SOC")
    titles: Dict[str, str] = {}
    to_soc: Dict[str, Set[str]] = defaultdict(set)
    for i in range(sheet.nrows):
        v = sheet.row_values(i)
        code = str(v[0]).strip()
        if not re.fullmatch(r"\d{4}", code):
            continue
        titles[code] = str(v[1]).strip()
        soc = str(v[3]).strip()
        if re.fullmatch(r"\d\d-\d{4}", soc):
            to_soc[code].add(soc)
    return titles, to_soc


def load_soc2010_to_2018() -> Dict[str, Set[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(SOURCES / "soc2010_to_soc2018_crosswalk.xlsx", read_only=True)
    ws = wb.worksheets[0]
    out: Dict[str, Set[str]] = defaultdict(set)
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and r[2] and re.fullmatch(r"\d\d-\d{4}", str(r[0]).strip()):
            out[str(r[0]).strip()].add(str(r[2]).strip())
    return out


def load_careers():
    careers = []
    for p in sorted(CAREERS_DIR.glob("*.json")):
        d = json.loads(p.read_text(encoding="utf-8"))
        if len(d.get("skills", [])) < 5:
            continue  # not listed in the app
        soc = ((d.get("onet") or {}).get("soc_code") or "")[:7]
        careers.append({"slug": d["slug"], "name": d["name"], "soc": soc})
    return careers


def match_unit_to_isco(unit_title: str, unit_code: str, isco_titles: Dict[str, str]) -> Tuple[str, float]:
    best_code, best = "", 0.0
    for code, title in isco_titles.items():
        s = sim(unit_title, title)
        if code == unit_code:
            s += 0.05  # same number breaks ties
        if s > best:
            best_code, best = code, s
    return best_code, best


def build() -> dict:
    units, occupations, alt = load_ofo()
    isco_titles, isco_to_soc = load_isco_to_soc2010()
    soc10_to_18 = load_soc2010_to_2018()
    careers = load_careers()
    careers_by_soc: Dict[str, List[dict]] = defaultdict(list)
    for c in careers:
        careers_by_soc[c["soc"]].append(c)

    unit_to_isco = {}
    for code, title in units.items():
        isco_code, score = match_unit_to_isco(title, code, isco_titles)
        if score >= 0.5:
            unit_to_isco[code] = isco_code

    results, unmatched = [], []
    stats = defaultdict(int)
    for code, title in occupations:
        unit = code[:4]
        candidates: List[dict] = []
        isco = unit_to_isco.get(unit)
        if isco:
            for soc10 in isco_to_soc.get(isco, ()):
                for soc18 in soc10_to_18.get(soc10, {soc10}):
                    candidates += careers_by_soc.get(soc18, [])
        names = [title] + alt.get(code, [])

        def score(c):
            return max(overlap(n, c["name"]) + 0.5 * sim(n, c["name"]) for n in names)

        chosen, how = None, ""
        if candidates:
            uniq = {c["slug"]: c for c in candidates}.values()
            ranked = sorted(uniq, key=score, reverse=True)
            chosen, how = ranked[0], "crosswalk"
        else:
            ranked = sorted(careers, key=score, reverse=True)
            if ranked and score(ranked[0]) >= 1.0:
                chosen, how = ranked[0], "title"
        if not chosen:
            unmatched.append(f"{code} {title}")
            stats["unmatched"] += 1
            continue
        stats[how] += 1
        aliases = sorted({a for a in alt.get(code, []) if a.lower() != title.lower()})
        results.append({
            "ofo_code": code,
            "title": title,
            "aliases": aliases,
            "career_slug": chosen["slug"],
            "career_name": chosen["name"],
            "match": how,
        })

    return {
        "source": "OFO 2021 (South Africa) linked to O*NET careers via ISCO-08 and SOC crosswalks",
        "counts": dict(stats),
        "occupations": results,
        "unmatched": unmatched,
    }


def main():
    data = build()
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {OUT}: {data['counts']}")


if __name__ == "__main__":
    sys.exit(main())
