"""CV files in every format people actually send: Word templates with text
boxes, two-column layouts, old .doc files, WhatsApp downloads with no
extension, OpenDocument, RTF and plain text."""
import io
import os
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml

from app.extractors.cv_extractor import CVExtractionError, detect_format, extract_cv

FILES = os.path.join(os.path.dirname(__file__), "files")

LINES = [
    "Thandi Mokoena",
    "Experience",
    "Accounts Clerk, Mabena Trading, 2022 - 2025",
    "Reconciled 200+ supplier accounts monthly using Pastel and Excel.",
    "Education",
    "National Diploma: Financial Accounting, 2021",
]

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def _paras(lines):
    return "".join(f"<w:p><w:r><w:t>{l}</w:t></w:r></w:p>" for l in lines)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain_docx() -> bytes:
    d = Document()
    for l in LINES:
        d.add_paragraph(l)
    return _save(d)


def _textbox_docx() -> bytes:
    """Whole CV inside a text box, written the way Word writes it: a
    DrawingML shape plus a VML fallback copy of the same text."""
    d = Document()
    box = (
        f'<w:p {NS}><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing>'
        '<wp:inline><wp:extent cx="1" cy="1"/><wp:docPr id="1" name="Text Box 1"/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f"<wps:wsp><wps:txbx><w:txbxContent>{_paras(LINES)}</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp>"
        "</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{_paras(LINES)}</w:txbxContent>"
        "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>"
    )
    d.element.body.insert(0, parse_xml(box))
    return _save(d)


def _nested_table_docx() -> bytes:
    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Contact: Pretoria"
    inner = t.cell(0, 1).add_table(rows=len(LINES), cols=1)
    for i, l in enumerate(LINES):
        inner.cell(i, 0).text = l
    return _save(d)


def _odt() -> bytes:
    body = "".join(f"<text:p>{l}</text:p>" for l in LINES)
    content = (
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        f"<office:body><office:text>{body}</office:text></office:body></office:document-content>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", content)
    return buf.getvalue()


def _rtf() -> bytes:
    body = "\\par\n".join(LINES)
    return ("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs22 " + body + "}").encode("latin-1")


def _assert_full_cv(cv):
    assert "Thandi Mokoena" in cv.raw_text
    assert cv.raw_text.count("Reconciled 200+ supplier accounts") == 1  # no duplicate fallback text
    assert "Reconciled" in cv.sections["experience"]


@pytest.mark.parametrize(
    "name,make",
    [
        ("cv.docx", _plain_docx),
        ("template.docx", _textbox_docx),
        ("two-column.docx", _nested_table_docx),
        ("cv.odt", _odt),
        ("cv.rtf", _rtf),
        ("cv.txt", lambda: "\n".join(LINES).encode("utf-8")),
    ],
)
def test_reads_every_supported_format(name, make):
    _assert_full_cv(extract_cv(name, make()))


def test_old_word_doc():
    with open(os.path.join(FILES, "sample_cv.doc"), "rb") as f:
        data = f.read()
    cv = extract_cv("cv.doc", data)
    assert "Reconciled 200+ supplier accounts" in cv.raw_text
    assert "Reconciled" in cv.sections["experience"]


def test_old_doc_saved_with_docx_name():
    with open(os.path.join(FILES, "sample_cv.doc"), "rb") as f:
        data = f.read()
    assert detect_format("cv.docx", data) == "doc"
    assert "Reconciled" in extract_cv("cv.docx", data).raw_text


def test_whatsapp_download_with_no_extension():
    cv = extract_cv("DOC-20261030-WA0004", _plain_docx())
    _assert_full_cv(cv)


def test_photo_gets_a_helpful_message():
    jpeg = b"\xff\xd8\xff\xe0\x00\x10JFIF\x00" + b"\x00" * 200
    with pytest.raises(CVExtractionError, match="photo"):
        extract_cv("cv.jpg", jpeg)


def test_random_binary_rejected():
    with pytest.raises(CVExtractionError, match="can't read this type"):
        extract_cv("cv.bin", bytes(range(256)) * 20)
