"""
CV text extraction.

Supports PDF (via PyMuPDF) and DOCX (via python-docx).
Also performs light section detection (experience / education / projects /
skills) using heading heuristics, so downstream evidence analysis can use
section context (e.g. a skill mentioned inside "Experience" carries more
weight than one mentioned in an unrelated section).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Dict, List

import fitz  # PyMuPDF
from docx import Document


class CVExtractionError(Exception):
    """Raised when a CV file cannot be read or contains no usable text."""


SECTION_HEADINGS = {
    "experience": ["experience", "work experience", "employment history", "professional experience"],
    "education": ["education", "academic background", "qualifications"],
    "projects": ["projects", "personal projects", "portfolio"],
    "skills": ["skills", "technical skills", "core competencies", "technologies"],
    "summary": ["summary", "profile", "about", "objective"],
}


@dataclass
class ExtractedCV:
    raw_text: str
    sections: Dict[str, str] = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        return len(self.raw_text.strip()) < 30


def _detect_sections(lines: List[str]) -> Dict[str, str]:
    sections: Dict[str, List[str]] = {key: [] for key in SECTION_HEADINGS}
    sections["other"] = []
    current = "other"

    for line in lines:
        stripped = line.strip()
        lowered = stripped.lower().strip(":").strip()
        matched_section = None
        if 0 < len(lowered) <= 40:
            for key, headings in SECTION_HEADINGS.items():
                if lowered in headings:
                    matched_section = key
                    break
        if matched_section:
            current = matched_section
            continue
        sections[current].append(stripped)

    return {key: "\n".join(v) for key, v in sections.items()}


def extract_from_pdf(file_bytes: bytes) -> ExtractedCV:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise CVExtractionError(f"Could not open PDF: {exc}") from exc

    try:
        text_parts = []
        for page in doc:
            text_parts.append(page.get_text("text"))
        raw_text = "\n".join(text_parts)
    finally:
        doc.close()

    if not raw_text.strip():
        raise CVExtractionError(
            "No extractable text found in PDF. The file may be a scanned "
            "image without a text layer."
        )

    lines = [l for l in raw_text.splitlines()]
    return ExtractedCV(raw_text=raw_text, sections=_detect_sections(lines))


def extract_from_docx(file_bytes: bytes) -> ExtractedCV:
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001
        raise CVExtractionError(f"Could not open DOCX: {exc}") from exc

    lines = []

    # Extract text from paragraphs
    for p in document.paragraphs:
        if p.text.strip():
            lines.append(p.text)

    # Extract text from tables (many CVs use table layouts)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                # Each cell can contain multiple paragraphs
                cell_text = cell.text.strip()
                if cell_text:
                    lines.append(cell_text)

    raw_text = "\n".join(lines)

    if not raw_text.strip():
        raise CVExtractionError("No extractable text found in DOCX file.")

    return ExtractedCV(raw_text=raw_text, sections=_detect_sections(lines))


def extract_from_text(raw_text: str) -> ExtractedCV:
    """Build an ExtractedCV directly from plain text (used by tests and any
    future direct-text-paste feature)."""
    lines = raw_text.splitlines()
    return ExtractedCV(raw_text=raw_text, sections=_detect_sections(lines))


def extract_cv(filename: str, file_bytes: bytes) -> ExtractedCV:
    """Dispatch extraction based on file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_from_pdf(file_bytes)
    if lower.endswith(".docx"):
        return extract_from_docx(file_bytes)
    raise CVExtractionError(
        "Unsupported file format. Please upload a PDF or DOCX file."
    )


YEAR_PATTERN = re.compile(r"\b(19|20)\d{2}\b")


def extract_years_mentioned(text: str) -> List[int]:
    """Best-effort extraction of 4-digit years for recency scoring."""
    return sorted({int(m) for m in re.findall(r"\b((?:19|20)\d{2})\b", text)})
